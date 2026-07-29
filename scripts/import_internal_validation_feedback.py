"""Import human feedback from the completed campaign DOCX into reviewed JSON."""

from __future__ import annotations

import argparse
import json
import re
from pathlib import Path

from docx import Document


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("docx", type=Path)
    parser.add_argument("--campaign-json", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()
    campaign = json.loads(args.campaign_json.read_text(encoding="utf-8"))
    document = Document(args.docx)
    feedback_tables = [table for table in document.tables if table.cell(0, 0).text.strip() == "Outcome"]
    if len(feedback_tables) != len(campaign["cases"]):
        raise SystemExit(f"Tabelle feedback trovate: {len(feedback_tables)}; attese: {len(campaign['cases'])}")
    errors = []
    for case, table in zip(campaign["cases"], feedback_tables):
        values = {row.cells[0].text.strip(): row.cells[1].text.strip() for row in table.rows}
        outcome = normalize_outcome(values.get("Outcome", ""))
        rating = parse_score(values.get("Rating", ""))
        clarity = parse_score(values.get("Chiarezza", ""))
        usefulness = parse_score(values.get("Utilità", ""))
        notes = values.get("Note", "").strip()
        if not outcome or rating is None or clarity is None or usefulness is None:
            errors.append(case["review_id"])
            continue
        case["human_feedback"] = {
            "outcome": outcome,
            "rating": rating,
            "clarity": clarity,
            "usefulness": usefulness,
            "notes": "" if notes.startswith("[COMPILARE") else notes[:1000],
        }
    if errors:
        raise SystemExit("Feedback incompleto o non valido per: " + ", ".join(errors))
    campaign["human_review_status"] = "completed"
    campaign["human_summary"] = summarize(campaign["cases"])
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(campaign, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(campaign["human_summary"], ensure_ascii=False))
    return 0


def normalize_outcome(value: str) -> str | None:
    normalized = value.lower().strip()
    mapping = {"corretto": "correct", "parziale": "partial", "errato": "incorrect",
               "correct": "correct", "partial": "partial", "incorrect": "incorrect"}
    return mapping.get(normalized)


def parse_score(value: str) -> int | None:
    match = re.fullmatch(r"\s*([1-5])\s*", value)
    return int(match.group(1)) if match else None


def summarize(cases: list[dict]) -> dict:
    outcomes = {"correct": 0, "partial": 0, "incorrect": 0}
    ratings = []
    for case in cases:
        feedback = case["human_feedback"]
        outcomes[feedback["outcome"]] += 1
        ratings.append(feedback["rating"])
    return {"total": len(cases), "outcomes": outcomes, "average_rating": round(sum(ratings) / len(ratings), 2)}


if __name__ == "__main__":
    raise SystemExit(main())
